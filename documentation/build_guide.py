from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).parent
SOURCE = ROOT / 'guide_content.md'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.5)
sec.left_margin = sec.right_margin = Inches(.75)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s = doc.styles[name]
    s.font.name = 'Century Gothic'
    s.font.color.rgb = RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Century Gothic')
    s.element.rPr.rFonts.set(qn('w:hAnsi'), 'Century Gothic')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        s.element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.0
for name, size in [('Title',18),('Subtitle',12),('Heading 1',14),('Heading 2',12)]:
    s=doc.styles[name]; s.font.size=Pt(size)
    s.paragraph_format.space_before=Pt(7 if name=='Heading 2' else 0)
    s.paragraph_format.space_after=Pt(5)
    s.paragraph_format.keep_with_next=True
doc.core_properties.title='DefendIt System Guide and Documentation'
doc.core_properties.subject='Thesis Consultation and Milestone Tracker'
doc.core_properties.author='DefendIt Project Team'
foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=foot.add_run('DefendIt  |  '); run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)

def table(lines):
    records=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    records=[r for r in records if not all(re.fullmatch(r'[-: ]+',v) for v in r)]
    t=doc.add_table(rows=0,cols=len(records[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths={2:[1.8,5.2],3:[1.5,1.1,4.4],4:[.62,1.1,2.55,2.73],5:[1.1,1.48,1.48,1.48,1.46]}[len(records[0])]
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,rec in enumerate(records):
        row=t.add_row()
        trpr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trpr.append(no)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for cell,value,w in zip(row.cells,rec,widths):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=value
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EAF0' if i==0 else 'FFFFFF');pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'45');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.size=Pt(11);r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

content=SOURCE.read_text(encoding='utf-8')
wire=Image.new('RGB',(1800,660),'white')
draw=ImageDraw.Draw(wire)
font=ImageFont.truetype('C:/Windows/Fonts/GOTHIC.TTF',29)
bold=ImageFont.truetype('C:/Windows/Fonts/GOTHICB.TTF',32)
frames=[('Today',['Greeting     Sync status','Next consultation','Summary counts','Needs attention','Open action items','Today  Schedule  Groups  Progress']),('Group and titles',['Group name and adviser','Section and roster','Title ideas and status','Proposal details','Add or review title','Today  Schedule  Groups  Progress']),('Milestone detail',['Milestone and status','Due date and progress','Tasks and assignees','Team workload','Submit or review','Today  Schedule  Groups  Progress'])]
for idx,(title,labels) in enumerate(frames):
    x=idx*600+12
    draw.rounded_rectangle((x,8,x+570,646),radius=15,outline='black',width=3)
    draw.text((x+20,25),title,font=bold,fill='black')
    for j,label in enumerate(labels):
        y=95+j*85
        draw.rectangle((x+18,y,x+550,y+65),outline='#777777',width=2)
        f=font if j<5 else ImageFont.truetype('C:/Windows/Fonts/GOTHIC.TTF',21)
        draw.text((x+29,y+13),label,font=f,fill='black')
wire.save(ROOT/'wireframes.png')
for page_no,page in enumerate(content.split('\n===PAGE===\n')):
    if page_no:doc.add_page_break()
    lines=page.strip().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i]);i+=1
            table(block);continue
        if line=='[WIREFRAMES]':doc.add_picture(str(ROOT/'wireframes.png'),width=Inches(7))
        elif line.startswith('# '):doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('- '):
            p=doc.add_paragraph(line[2:]);p.paragraph_format.left_indent=Inches(.15)
            p.text='• '+p.text
        else:doc.add_paragraph(line)
        i+=1
out=ROOT/'DefendIt_System_Guide.docx'
doc.save(out)
print(out)
print('Word count:',len(re.findall(r"\b[\w]+(?:[’'-][\w]+)*\b",content)))
